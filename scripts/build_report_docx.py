"""Turn docs/PROJECT_REPORT.md and docs/STUDENT_GUIDE.md (made by build_report.py) into .docx files with the figures embedded.

  python scripts/build_report.py && python scripts/build_report_docx.py

Handles only the markdown those two files use: headings, paragraphs, bullet and numbered lists, block quotes, pipe tables,
images, code blocks, **bold**, *italic* and `code`. Needs only python-docx (no torch, no network). Nothing is typed by hand:
the text is whatever the markdown says, so the numbers still come from results/.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
JOBS = [("PROJECT_REPORT.md", "PROJECT_REPORT.docx", "Appendix B"), ("STUDENT_GUIDE.md", "STUDENT_GUIDE.docx", None)]  # heading that starts a landscape section
BODY, MONO, NAVY = "Calibri", "Consolas", RGBColor(0x1F, 0x38, 0x64)
INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*\s](?:[^*]*[^*\s])?\*|`[^`]+`)")
NUMERIC = re.compile(r"^[\d,.+\-−x() %]+$")


def shade(element, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"), shd.set(qn("w:color"), "auto"), shd.set(qn("w:fill"), fill)
    element.append(shd)


def runs(par, text: str, size: float | None = None, bold: bool = False) -> None:
    """Add `text` to a paragraph, turning **bold**, *italic* and `code` into formatted runs."""
    for part in INLINE.split(text):
        if not part:
            continue
        code, strong, emph = part.startswith("`"), part.startswith("**") and len(part) > 4, part.startswith("*") and not part.startswith("**")
        run = par.add_run(part[2:-2] if strong else part[1:-1] if code or emph else part)
        run.bold = bold or strong or None
        run.italic = emph or None
        if code:
            run.font.name, run.font.size = MONO, Pt((size or 11) - 1)
            shade(run._r.get_or_add_rPr(), "F2F2F2")
        elif size:
            run.font.size = Pt(size)


def new_document(title: str) -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2.2)
    styles = doc.styles
    styles["Normal"].font.name, styles["Normal"].font.size = BODY, Pt(11)
    styles["Normal"].paragraph_format.space_after, styles["Normal"].paragraph_format.line_spacing = Pt(6), 1.15
    for name, size in (("Title", 22), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        st = styles[name]
        st.font.name, st.font.size, st.font.bold, st.font.color.rgb = BODY, Pt(size), True, NAVY
        st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(0 if name == "Title" else 14), Pt(6)
    styles["Caption"].font.name, styles["Caption"].font.size, styles["Caption"].font.italic = BODY, Pt(9), True
    styles["Caption"].font.color.rgb = RGBColor(0x40, 0x40, 0x40)
    foot = sec.footer.paragraphs[0]                                             # page number field
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for kind in ("begin", "PAGE", "separate", "1", "end"):
        run = foot.add_run()
        if kind == "PAGE":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = kind
            run._r.append(el)
        elif kind == "1":
            run.text = kind
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            run._r.append(el)
    doc.core_properties.title, doc.core_properties.author, doc.core_properties.last_modified_by = title, "", ""
    return doc


def text_width(doc: Document) -> float:
    sec = doc.sections[-1]
    return (sec.page_width - sec.left_margin - sec.right_margin) / 360000          # cm


def add_table(doc: Document, rows: list) -> None:
    header, body = rows[0], rows[1:]
    ncol = len(header)
    size = 9 if ncol <= 4 else 8.5 if ncol <= 8 else 7.5
    weights = [min(45, max(12, 2 + max([len(r[c]) for r in body] + [2 + max(len(w) for w in header[c].split())]))) for c in range(ncol)]
    table = doc.add_table(rows=len(rows), cols=ncol)
    table.style, table.alignment, table.autofit = "Table Grid", WD_TABLE_ALIGNMENT.CENTER, False
    center = [all(NUMERIC.match(r[c]) for r in body) for c in range(ncol)]
    for c in range(ncol):
        table.columns[c].width = Cm(text_width(doc) * weights[c] / sum(weights))
    for i, row in enumerate(rows):
        tr = table.rows[i]._tr
        trpr = tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if i == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        for c, cell_text in enumerate(row):
            cell = table.cell(i, c)
            cell.width = Cm(text_width(doc) * weights[c] / sum(weights))
            par = cell.paragraphs[0]
            par.paragraph_format.space_after, par.paragraph_format.line_spacing = Pt(1), 1.0
            if i == 0:
                shade(cell._tc.get_or_add_tcPr(), "D9E2F3")
            if center[c] or i == 0:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            runs(par, cell_text, size, bold=i == 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_image(doc: Document, alt: str, path: Path) -> None:
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.keep_with_next = True
    par.add_run().add_picture(str(path), width=Cm(min(16.0, text_width(doc))))
    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    m = re.match(r"(Figure \d+)(.*)", alt)
    if m:
        cap.add_run(m.group(1)).bold = True
        cap.add_run(m.group(2))
    else:
        cap.add_run(alt)


def add_quote(doc: Document, text: str) -> None:
    par = doc.add_paragraph()
    ppr = par._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    for k, v in (("val", "single"), ("sz", "18"), ("space", "8"), ("color", "8EA9DB")):
        left.set(qn(f"w:{k}"), v)
    border.append(left)
    ppr.append(border)
    shade(ppr, "F3F6FB")
    par.paragraph_format.left_indent = Cm(0.6)
    runs(par, text)


def add_code(doc: Document, lines: list) -> None:
    par = doc.add_paragraph()
    shade(par._p.get_or_add_pPr(), "F2F2F2")
    par.paragraph_format.left_indent, par.paragraph_format.line_spacing = Cm(0.5), 1.0
    for k, line in enumerate(lines):
        run = par.add_run(line)
        run.font.name, run.font.size = MONO, Pt(9)
        if k < len(lines) - 1:
            run.add_break()


def add_numbered(doc: Document, number: str, text: str) -> None:
    """A numbered item typed out, so each list starts at 1 and the references keep their numbers."""
    par = doc.add_paragraph()
    par.paragraph_format.left_indent, par.paragraph_format.first_line_indent = Cm(0.9), Cm(-0.9)
    par.paragraph_format.space_after = Pt(4)
    par.add_run(number + "\t")
    par.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))
    runs(par, text)


def convert(src: Path, dst: Path, landscape_from: str | None) -> None:
    lines = src.read_text(encoding="utf-8").splitlines()
    title = next(l[2:] for l in lines if l.startswith("# "))
    doc = new_document(title)
    centered, i = 0, 0
    while i < len(lines):
        line = lines[i]
        if not line.strip() or line.startswith("<!--"):
            i += 1
            continue
        if line.startswith("```"):
            j = i + 1
            while not lines[j].startswith("```"):
                j += 1
            add_code(doc, lines[i + 1:j])
            i = j + 1
            continue
        m = re.match(r"(#{1,3}) (.*)", line)
        if m:
            level, text = len(m.group(1)), m.group(2)
            if landscape_from and text.startswith(landscape_from):
                sec = doc.add_section(WD_SECTION.NEW_PAGE)
                sec.orientation = WD_ORIENT.LANDSCAPE
                sec.page_width, sec.page_height = sec.page_height, sec.page_width
            par = doc.add_paragraph(style="Title" if level == 1 else f"Heading {level - 1}")
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else None
            par.add_run(text)
            centered = 2 if level == 1 else 0
            i += 1
            continue
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            add_table(doc, [r for r in block if not all(re.fullmatch(r":?-{3,}:?", c) for c in r)])
            continue
        m = re.match(r"!\[(.*)\]\((.*)\)\s*$", line)
        if m:
            add_image(doc, m.group(1), (src.parent / m.group(2)).resolve())
        elif line.startswith("> "):
            add_quote(doc, line[2:])
        elif line.startswith("- "):
            runs(doc.add_paragraph(style="List Bullet"), line[2:])
        elif re.match(r"\d+\. ", line):
            number, text = line.split(". ", 1)
            add_numbered(doc, number + ".", text)
        else:
            par = doc.add_paragraph()
            if centered and len(line) < 250:
                par.alignment, centered = WD_ALIGN_PARAGRAPH.CENTER, centered - 1
            else:
                centered = 0
            runs(par, line)
        i += 1
    doc.save(dst)


def main() -> int:
    for md, out, landscape_from in JOBS:
        convert(DOCS / md, DOCS / out, landscape_from)
        print("wrote", (DOCS / out).relative_to(ROOT))
    return 0


if __name__ == "__main__":
    sys.exit(main())
